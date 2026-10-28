from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document as WordDocument
from docx.shared import Inches
import pytest
from app.core.config import get_settings
from app.rag.parsers import _ocr_image, parse_document


def test_markdown_heading_and_utf8(tmp_path: Path):
    path = tmp_path / "course.md"
    path.write_text("# 第二章 进程\n\n进程是运行中的程序实例。", encoding="utf-8")
    segments = parse_document(path, "MD")
    assert len(segments) == 1
    assert segments[0].section == "第二章 进程"
    assert "运行中的程序实例" in segments[0].text


def test_docx_preserves_paragraph_table_order(tmp_path: Path):
    path = tmp_path / "course.docx"
    doc = WordDocument()
    doc.add_heading("第四章 并发", level=1)
    doc.add_paragraph("先出现的段落")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "临界区"
    table.cell(0, 1).text = "互斥"
    doc.add_paragraph("后出现的段落")
    doc.save(path)
    segments = parse_document(path, "DOCX")
    assert [segment.text for segment in segments] == ["先出现的段落", "临界区 | 互斥", "后出现的段落"]
    assert all(segment.section == "第四章 并发" for segment in segments)


def test_pdf_pages_are_preserved(tmp_path: Path):
    path = tmp_path / "course.pdf"
    pdf = fitz.open()
    for text in ("First page course material.", "Second page course material."):
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()
    segments = parse_document(path, "PDF")
    assert [(item.page, item.text) for item in segments] == [(1, "First page course material."), (2, "Second page course material.")]


def test_ocr_loads_models_from_a_directory_with_spaces(monkeypatch):
    settings = get_settings()
    if not settings.ocr_enabled or not Path(settings.tesseract_cmd).is_file():
        pytest.skip("Tesseract OCR is not enabled in this environment")
    model_source = Path(settings.tessdata_dir)
    required = [model_source / "chi_sim.traineddata", model_source / "eng.traineddata"]
    if not all(path.is_file() for path in required):
        pytest.skip("OCR language data is not installed in this environment")
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if not font_path.is_file():
        pytest.skip("OCR sample font is not installed")
    temp_root = Path("D:/Personal/Temp")
    if not temp_root.is_dir():
        pytest.skip("No ASCII temporary directory is available for the Windows Tesseract smoke test")

    with TemporaryDirectory(prefix="studypath-ocr-", dir=temp_root) as temp_dir:
        model_dir = Path(temp_dir) / "OCR model data"
        model_dir.mkdir()
        for source in required:
            shutil.copyfile(source, model_dir / source.name)
        monkeypatch.setattr(settings, "tessdata_dir", str(model_dir))
        monkeypatch.setenv("TESSDATA_PREFIX", str(model_dir))

        image = Image.new("RGB", (1000, 180), "white")
        ImageDraw.Draw(image).text((30, 35), "StudyPath OCR READY", font=ImageFont.truetype(str(font_path), 64), fill="black")
        from io import BytesIO
        image_bytes = BytesIO()
        image.save(image_bytes, format="PNG")

        recognized = _ocr_image(image_bytes.getvalue(), "runtime test")
        assert "StudyPath" in recognized
        assert "READY" in recognized


def test_scanned_pdf_and_docx_image_text_are_ocr_indexable(tmp_path: Path):
    settings = get_settings()
    if not settings.ocr_enabled or not Path(settings.tesseract_cmd).is_file():
        pytest.skip("Tesseract OCR is not enabled in this environment")
    if not all((Path(settings.tessdata_dir) / name).is_file() for name in ("chi_sim.traineddata", "eng.traineddata")):
        pytest.skip("OCR language data is not installed in this environment")
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if not font_path.is_file():
        pytest.skip("OCR sample font is not installed")

    from io import BytesIO
    image = Image.new("RGB", (1200, 220), "white")
    ImageDraw.Draw(image).text((30, 55), "StudyPath OCR READY", font=ImageFont.truetype(str(font_path), 72), fill="black")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    raw_image = image_bytes.getvalue()

    pdf_path = tmp_path / "scan.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=600, height=110)
    page.insert_image(page.rect, stream=raw_image)
    pdf.save(pdf_path)
    pdf.close()
    pdf_segments = parse_document(pdf_path, "PDF")
    assert any("StudyPath" in segment.text and "READY" in segment.text and segment.page == 1 for segment in pdf_segments)

    docx_path = tmp_path / "embedded-image.docx"
    word = WordDocument()
    word.add_paragraph("Document image:").add_run().add_picture(BytesIO(raw_image), width=Inches(6))
    word.save(docx_path)
    docx_segments = parse_document(docx_path, "DOCX")
    assert any("StudyPath" in segment.text and "READY" in segment.text for segment in docx_segments)
