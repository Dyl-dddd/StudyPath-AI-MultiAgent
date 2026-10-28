from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
files = [
    ("00_登记材料要求与清单.md", "00_登记材料要求与清单.docx"),
    ("01_软件基本信息采集表.md", "01_软件基本信息采集表.docx"),
    ("02_软件说明书与用户手册草案.md", "02_软件说明书与用户手册_V1.2.docx"),
]

def set_font(run, size=10.5, bold=None, color=None, name='Microsoft YaHei'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def add_page_field(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)

for source, target in files:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.75), Cm(1.75)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.0)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.1
    settings = doc.settings.element
    zoom = settings.find(qn('w:zoom'))
    if zoom is not None:
        zoom.set(qn('w:percent'), '100')
    for style_name, size in [('Title', 20), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
        style = doc.styles[style_name]
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style_name == 'Title' and style._element.pPr is not None:
            border = style._element.pPr.find(qn('w:pBdr'))
            if border is not None:
                style._element.pPr.remove(border)
    header = section.header.paragraphs[0]
    header.alignment = 2
    set_font(header.add_run('智学研途 AI 个性化学习平台（StudyPath AI）V1.0'), 8, color=(100, 100, 100))
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    set_font(footer.add_run('著作权人：戴云龙　·　申报前请复核　·　第'), 8, color=(130, 125, 145))
    add_page_field(footer)
    set_font(footer.add_run('页'), 8, color=(130, 125, 145))

    lines = (root / source).read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        image_match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image_match:
            label, image_path = image_match.groups()
            image_file = root / image_path
            if not image_file.exists():
                raise FileNotFoundError(f'截图文件不存在：{image_file}')
            image_paragraph = doc.add_paragraph()
            image_paragraph.alignment = 1
            image_paragraph.paragraph_format.space_before = Pt(3)
            image_paragraph.paragraph_format.space_after = Pt(2)
            image_paragraph.paragraph_format.keep_with_next = True
            image_paragraph.add_run().add_picture(str(image_file), width=Cm(16.0))
            caption = doc.add_paragraph()
            caption.alignment = 1
            caption.paragraph_format.space_after = Pt(8)
            set_font(caption.add_run(label), 9, color=(100, 93, 120))
            i += 1; continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip()); i += 1
            rows = []
            for item in table_lines:
                cells = [c.strip() for c in item.strip('|').split('|')]
                if all(re.fullmatch(r':?-{2,}:?', c or '-') for c in cells): continue
                rows.append(cells)
            if rows:
                cols = max(map(len, rows))
                table = doc.add_table(rows=0, cols=cols)
                table.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    cells = table.add_row().cells
                    tr_pr = table.rows[-1]._tr.get_or_add_trPr()
                    no_split = OxmlElement('w:cantSplit')
                    tr_pr.append(no_split)
                    if ri == 0:
                        repeat_header = OxmlElement('w:tblHeader')
                        repeat_header.set(qn('w:val'), 'true')
                        tr_pr.append(repeat_header)
                    for ci in range(cols):
                        val = row[ci] if ci < len(row) else ''
                        p = cells[ci].paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        r = p.add_run(val)
                        set_font(r, 9, bold=(ri == 0), color=(76,55,145) if ri == 0 else None)
            continue
        if line.startswith('```'):
            i += 1; code=[]
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.3)
            r = p.add_run('\n'.join(code)); set_font(r, 9, name='Consolas')
            i += 1; continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title' if i < 3 else 'Heading 1')
            p.paragraph_format.keep_with_next = True
            set_font(p.add_run(line[2:].replace('**','')), 20 if i < 3 else 16, True, (0,0,0))
        elif line.startswith('## '):
            p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.keep_with_next=True; set_font(p.add_run(line[3:].replace('**','')),16,True,(0,0,0))
        elif line.startswith('### '):
            p=doc.add_paragraph(style='Heading 2'); p.paragraph_format.keep_with_next=True; set_font(p.add_run(line[4:].replace('**','')),13,True,(0,0,0))
        elif line.startswith('> '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.4)
            set_font(p.add_run(line[2:].replace('**','')),9.5,color=(116,108,135))
        elif re.match(r'^\d+\.\s', line):
            p=doc.add_paragraph(); set_font(p.add_run(line.replace('**','').replace('`','')),10.5)
        elif line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); set_font(p.add_run(line[2:].replace('**','').replace('`','')),10.5)
        else:
            p=doc.add_paragraph()
            text = line.replace('**','').replace('`','')
            if line.startswith('**图 '): p.paragraph_format.keep_with_next = True
            set_font(p.add_run(text), 10.5)
        i += 1
    doc.save(root / target)
    print(target)
