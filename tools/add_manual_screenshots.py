from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

source_path = r'D:\数字媒体竞赛\docs\copyright\02_软件说明书与用户手册_修订版.docx'
dest_path = r'D:\数字媒体竞赛\docs\copyright\02_软件说明书与用户手册_修订版_补充截图.docx'
doc = Document(source_path)

def find_paragraph(prefix):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise ValueError(f'Anchor not found: {prefix}')

def insert_after(anchor, text='', style_source=None, picture=None, width=None, center=False, keep=False):
    new_xml = OxmlElement('w:p')
    anchor._p.addnext(new_xml)
    p = Paragraph(new_xml, anchor._parent)
    p.style = doc.styles['Normal']
    if style_source is not None:
        p.paragraph_format.space_before = style_source.paragraph_format.space_before
        p.paragraph_format.space_after = style_source.paragraph_format.space_after
        p.paragraph_format.line_spacing = style_source.paragraph_format.line_spacing
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if keep:
        p.paragraph_format.keep_with_next = True
    if picture:
        p.add_run().add_picture(picture, width=width)
    if text:
        run = p.add_run(text)
        if style_source is not None and style_source.runs:
            src = style_source.runs[0]
            if src._r.rPr is not None:
                run._r.insert(0, deepcopy(src._r.rPr))
    return p

caption_template = find_paragraph('图 1')
body_template = find_paragraph('首页是学生进入系统后的学习驾驶舱')
max_width = min(6.35, (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin) / 914400 - 0.12)

# Screenshot 3: updated learner overview after study activity.
anchor = find_paragraph('首页下部提供课程知识图谱预览')
p = insert_after(anchor, '图 3-1　首页学习情况：任务完成状态、课程进度与复习入口', caption_template, keep=True)
p = insert_after(p, picture=r'D:\Personal\Temp\codex-clipboard-b7b49f49-d07c-4dc4-917c-f20005a64d6a.png', width=Inches(max_width), center=True, keep=True)
p = insert_after(p, '首页学习情况示例', caption_template)
p = insert_after(p, '该截图展示账号已有学习记录后的首页状态：今日学习目标显示已完成 2 项、共 5 项，任务列表用勾选标出已完成任务；页面同时保留课程学习进度、每日测验趋势、课程知识图谱预览和近期学习计划。学生可以从首页了解当天完成情况，并继续未完成任务或进入课程、图谱和计划页面。各项数值随当前账号的实际记录变化。', body_template)

# Screenshot 1: selected single knowledge node in the course graph.
anchor = find_paragraph('若图谱显示 0 个节点')
p = insert_after(anchor, '图 7-1　知识点详情：掌握度、前置知识与学习操作', caption_template, keep=True)
p = insert_after(p, picture=r'D:\Personal\Temp\codex-clipboard-9bd0750a-81f0-42ac-bcfa-abfaf573603d.png', width=Inches(max_width), center=True, keep=True)
p = insert_after(p, '图7-1 单个知识节点详情', caption_template)
p = insert_after(p, '点击图谱中的单个节点后，右侧详情面板会展示知识点名称、所属课程、个人掌握度、知识点解释、前置知识和相关知识点。学生可在此标记掌握状态、跳转查看前置知识，并进入资料库、AI 讲解、学习计划或针对该知识点的练习。示例选中“处理器调度”，面板显示其掌握度为 0%，并关联“进程管理”；该数值仅表示当前账号记录。', body_template)

# Screenshot 2: analytics after saved quiz and task activity.
anchor = find_paragraph('推荐的复盘方式是：')
p = insert_after(anchor, '图 15　学习数据分析：测验记录、复习状态与待巩固知识点', caption_template, keep=True)
p = insert_after(p, picture=r'D:\Personal\Temp\codex-clipboard-92eff0c8-33e4-40b7-beaa-b98fafc55315.png', width=Inches(max_width), center=True, keep=True)
p = insert_after(p, '图15 学习活动后的数据分析示例', caption_template)
p = insert_after(p, '截图展示账号完成学习活动后形成的数据分析页面：顶部汇总所选周期内的测验次数、学习任务完成数、学习时长和平均正确率；中部用趋势图与环形图呈现测验及错题复习状态；下方按错题次数列出需要关注的知识点，并提供进入错题本复习的入口。示例中记录了 2 次测验和 2 项已完成任务，待复习知识点包括“进程同步与互斥”和“虚拟内存与分页”。页面统计以当前账号的已保存记录为准。', body_template)

figure_table = doc.tables[2]
def add_index_row(after_label, values):
    target = next(row for row in figure_table.rows if row.cells[0].text == after_label)
    row = figure_table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9)
    target._tr.addnext(row._tr)

add_index_row('图 3', ['图 3-1', '首页学习情况', '7.1', '任务完成状态、课程进度和复习入口'])
add_index_row('图 7', ['图 7-1', '单个知识节点详情', '7.4', '掌握度、节点解释、前置知识及学习操作'])
add_index_row('图 14', ['图 15', '学习数据分析示例', '7.10', '测验记录、复习状态和待巩固知识点'])

# Keep the numbered figure sequence aligned with the body order.
def replace_paragraph_prefix(prefix, old, new):
    paragraph = find_paragraph(prefix)
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)

replace_paragraph_prefix('图 15　课程内容维护', '图 15', '图 16')
replace_paragraph_prefix('图15 课程维护主界面', '图15', '图16')
replace_paragraph_prefix('图 16　课程内容维护下部', '图 16', '图 17')
replace_paragraph_prefix('图16 先修关系与题目维护', '图16', '图17')
for row in figure_table.rows:
    label, subject = row.cells[0].text, row.cells[1].text
    if label == '图 15' and subject == '课程维护主界面':
        row.cells[0].text = '图 16'
    elif label == '图 16' and subject == '先修关系与题目维护':
        row.cells[0].text = '图 17'
doc.save(dest_path)
print(f'Saved {dest_path}; paragraphs={len(doc.paragraphs)}; inline_shapes={len(doc.inline_shapes)}; image_width_in={max_width:.2f}')






